from pathlib import Path

from openpyxl import load_workbook


def load_excel(file_path: str) -> list[dict]:
    """Load financial records from an Excel (.xlsx) file."""

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx Excel files are supported.")

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        workbook.close()
        raise ValueError("Excel file is empty.")

    headers = rows[0]

    if not any(header is not None for header in headers):
        workbook.close()
        raise ValueError("Excel file does not contain a header row.")

    headers = [str(header).strip() if header is not None else "" for header in headers]

    records = []

    for row in rows[1:]:
        if all(value is None for value in row):
            continue

        record = {
            headers[index]: row[index]
            for index in range(min(len(headers), len(row)))
            if headers[index]
        }

        records.append(record)

    workbook.close()
    return records
