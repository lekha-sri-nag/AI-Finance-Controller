from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from models.invoice import Invoice
from models.exception import ExceptionRecord
from models.risk_result import RiskResult
from models.recommendation import Recommendation
from models.decision import Decision


REPORT_DIR = Path("data/reports")
REPORT_FILE = REPORT_DIR / "financial_analysis_report.xlsx"


def _format_sheet(sheet) -> None:
    """Apply basic formatting and readable column widths."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    for cell in sheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for column_cells in sheet.columns:
        max_length = 0
        column_index = column_cells[0].column

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(max_length + 2, 50)

    sheet.freeze_panes = "A2"


def _get_or_create_workbook():
    """Load the existing report or create a new workbook."""

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    if REPORT_FILE.exists():
        workbook = load_workbook(REPORT_FILE)

        required_sheets = [
            "Invoice Analysis",
            "Human Decisions",
            "Audit Trail"
        ]

        for sheet_name in required_sheets:
            if sheet_name not in workbook.sheetnames:
                workbook.create_sheet(title=sheet_name)

        return workbook

    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(default_sheet)

    workbook.create_sheet(title="Invoice Analysis")
    workbook.create_sheet(title="Human Decisions")
    workbook.create_sheet(title="Audit Trail")

    return workbook


def _initialize_sheets(workbook) -> None:
    """Create headers if the sheets are empty."""

    invoice_sheet = workbook["Invoice Analysis"]

    if invoice_sheet.max_row == 1 and invoice_sheet["A1"].value is None:
        invoice_sheet.append([
            "Invoice ID",
            "Vendor ID",
            "Purchase Order ID",
            "Invoice Date",
            "Amount",
            "Quantity",
            "Currency",
            "Invoice Status",
            "Exception Count",
            "Exceptions",
            "Risk Score",
            "Risk Level",
            "Risk Factors",
            "Risk Explanation",
            "AI Recommendation",
            "Recommendation Priority",
            "Human Review Required",
            "Recommendation Reason",
        ])

    decision_sheet = workbook["Human Decisions"]

    if decision_sheet.max_row == 1 and decision_sheet["A1"].value is None:
        decision_sheet.append([
            "Decision ID",
            "Invoice ID",
            "Recommendation ID",
            "AI Recommendation",
            "Human Decision",
            "Decision Maker",
            "Decision Time",
            "Comments",
            "Decision vs AI",
        ])

    audit_sheet = workbook["Audit Trail"]

    if audit_sheet.max_row == 1 and audit_sheet["A1"].value is None:
        audit_sheet.append([
            "Invoice ID",
            "Event",
            "Details",
        ])


def _existing_invoice_ids(sheet) -> set[str]:
    """Return invoice IDs already stored in the report."""

    invoice_ids = set()

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        if row[0]:
            invoice_ids.add(str(row[0]))

    return invoice_ids


def _get_ai_recommendation_from_sheet(
    invoice_sheet,
    invoice_id: str
) -> str:
    """
    Find the previously stored AI recommendation
    for an invoice.

    AI Recommendation is column O (15).
    """

    for row in invoice_sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        if not row:
            continue

        stored_invoice_id = row[0]

        if (
            stored_invoice_id
            and str(stored_invoice_id) == str(invoice_id)
        ):
            if len(row) >= 15 and row[14]:
                return str(row[14])

            return ""

    return ""


def create_excel_report(
    invoices: Iterable[Invoice],
    exceptions: Iterable[ExceptionRecord],
    risk_results: Iterable[RiskResult],
    recommendations: Iterable[Recommendation],
    decisions: Iterable[Decision] | None = None,
) -> str:
    """
    Create or update the Excel financial analysis report.

    Existing invoices are preserved.
    New invoices are appended.
    Duplicate invoice IDs are not added again.

    Human decisions can be saved independently after
    invoice processing.

    Returns:
        Path of the generated Excel file.
    """

    workbook = _get_or_create_workbook()

    _initialize_sheets(workbook)

    invoices = list(invoices)
    exceptions = list(exceptions)
    risk_results = list(risk_results)
    recommendations = list(recommendations)
    decisions = list(decisions or [])

    invoice_sheet = workbook["Invoice Analysis"]
    decision_sheet = workbook["Human Decisions"]
    audit_sheet = workbook["Audit Trail"]

    # ---------------------------------------------------------
    # Build lookup maps
    # ---------------------------------------------------------

    exception_map = {}

    for exception in exceptions:
        exception_map.setdefault(
            exception.invoice_id,
            []
        ).append(exception)

    risk_map = {
        risk.invoice_id: risk
        for risk in risk_results
    }

    recommendation_map = {
        recommendation.invoice_id: recommendation
        for recommendation in recommendations
    }

    # ---------------------------------------------------------
    # SHEET 1: Invoice Analysis
    # ---------------------------------------------------------

    existing_invoice_ids = _existing_invoice_ids(
        invoice_sheet
    )

    for invoice in invoices:

        # Do not duplicate an existing invoice.
        if invoice.invoice_id in existing_invoice_ids:
            continue

        invoice_exceptions = exception_map.get(
            invoice.invoice_id,
            []
        )

        risk = risk_map.get(
            invoice.invoice_id
        )

        recommendation = recommendation_map.get(
            invoice.invoice_id
        )

        invoice_sheet.append([
            invoice.invoice_id,
            invoice.vendor_id,
            invoice.purchase_order_id,
            invoice.invoice_date,
            invoice.amount,
            invoice.quantity,
            invoice.currency,
            invoice.status,
            len(invoice_exceptions),
            "; ".join(
                exception.exception_type
                for exception in invoice_exceptions
            ),
            risk.risk_score if risk else "",
            risk.risk_level if risk else "",
            "; ".join(
                risk.risk_factors
            ) if risk else "",
            risk.explanation if risk else "",
            recommendation.action
            if recommendation else "",
            recommendation.priority
            if recommendation else "",
            recommendation.requires_human_review
            if recommendation else "",
            recommendation.reason
            if recommendation else "",
        ])

        existing_invoice_ids.add(
            invoice.invoice_id
        )

    # ---------------------------------------------------------
    # SHEET 2: Human Decisions
    # ---------------------------------------------------------

    existing_decision_ids = set()

    for row in decision_sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        if row[0]:
            existing_decision_ids.add(
                str(row[0])
            )

    for decision in decisions:

        # Prevent duplicate human decisions.
        if decision.decision_id in existing_decision_ids:
            continue

        # First try to get the recommendation from
        # the current processing request.
        recommendation = recommendation_map.get(
            decision.invoice_id
        )

        if recommendation:
            ai_action = recommendation.action
        else:
            # If this is a later human decision,
            # retrieve the previously stored recommendation
            # from Invoice Analysis.
            ai_action = _get_ai_recommendation_from_sheet(
                invoice_sheet,
                decision.invoice_id
            )

        decision_comparison = ""

        if ai_action:
            ai_normalized = ai_action.strip().lower()
            human_normalized = decision.decision.strip().lower()

            decision_mapping = {
                "block payment": "block",
                "approve payment": "approve",
                "hold payment": "hold",
                "reject payment": "reject",
            }

            ai_normalized = decision_mapping.get(
                ai_normalized,
                ai_normalized
            )

            human_normalized = decision_mapping.get(
                human_normalized,
                human_normalized
            )

            decision_comparison = (
                "Agreed"
                if ai_normalized == human_normalized
                else "Overridden"
            )

        decision_sheet.append([
            decision.decision_id,
            decision.invoice_id,
            decision.recommendation_id,
            ai_action,
            decision.decision,
            decision.decided_by,
            decision.decided_at,
            decision.comments or "",
            decision_comparison,
        ])

        existing_decision_ids.add(
            decision.decision_id
        )

    # ---------------------------------------------------------
    # SHEET 3: Audit Trail
    # ---------------------------------------------------------

    existing_audit_entries = set()

    for row in audit_sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        if len(row) >= 3:
            existing_audit_entries.add(
                (
                    str(row[0]),
                    str(row[1]),
                    str(row[2])
                )
            )

    # ---------------------------------------------------------
    # Audit entries related to invoice processing
    # ---------------------------------------------------------

    for invoice in invoices:

        invoice_exceptions = exception_map.get(
            invoice.invoice_id,
            []
        )

        risk = risk_map.get(
            invoice.invoice_id
        )

        recommendation = recommendation_map.get(
            invoice.invoice_id
        )

        audit_entries = []

        audit_entries.append([
            invoice.invoice_id,
            "Invoice Processed",
            f"Invoice amount: {invoice.amount} "
            f"{invoice.currency}"
        ])

        for exception in invoice_exceptions:
            audit_entries.append([
                invoice.invoice_id,
                "Exception Detected",
                f"{exception.exception_type} - "
                f"{exception.severity}: "
                f"{exception.description}"
            ])

        if risk:
            audit_entries.append([
                invoice.invoice_id,
                "Risk Assessment",
                f"Score: {risk.risk_score}, "
                f"Level: {risk.risk_level}"
            ])

        if recommendation:
            audit_entries.append([
                invoice.invoice_id,
                "AI Recommendation",
                f"{recommendation.action} - "
                f"{recommendation.priority}"
            ])

        for decision in decisions:

            if decision.invoice_id != invoice.invoice_id:
                continue

            audit_entries.append([
                invoice.invoice_id,
                "Human Decision",
                f"{decision.decision} by "
                f"{decision.decided_by}"
            ])

        for entry in audit_entries:

            entry_key = (
                str(entry[0]),
                str(entry[1]),
                str(entry[2])
            )

            if entry_key in existing_audit_entries:
                continue

            audit_sheet.append(entry)

            existing_audit_entries.add(
                entry_key
            )

    # ---------------------------------------------------------
    # Audit human decisions even when invoices=[]
    # ---------------------------------------------------------

    for decision in decisions:

        audit_entry = [
            decision.invoice_id,
            "Human Decision",
            f"{decision.decision} by "
            f"{decision.decided_by}"
        ]

        entry_key = (
            str(audit_entry[0]),
            str(audit_entry[1]),
            str(audit_entry[2])
        )

        if entry_key in existing_audit_entries:
            continue

        audit_sheet.append(audit_entry)

        existing_audit_entries.add(
            entry_key
        )

    # ---------------------------------------------------------
    # Formatting
    # ---------------------------------------------------------

    _format_sheet(invoice_sheet)
    _format_sheet(decision_sheet)
    _format_sheet(audit_sheet)

    # ---------------------------------------------------------
    # Save workbook
    # ---------------------------------------------------------

    workbook.save(REPORT_FILE)

    return str(REPORT_FILE)